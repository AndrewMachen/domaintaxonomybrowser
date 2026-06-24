"""Excel import/export, faithful to the original Node/ExcelJS behavior.

* Export emits one row per branch, preserving empty branches.
* Import validates the column structure, then processes every row inside the
  caller's single transaction: blank rows are skipped, per-row validation
  problems are recorded and skipped, and an unexpected error propagates so the
  caller rolls the whole import back (no partial corruption).
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .service import TaxonomyService
from .types import Level
from .validation import ValidationError, normalize_technical_owner, same_name, validate_create

SHEET = "Taxonomy"

COLUMNS = [
    "Domain Group Name",
    "Domain Group Description",
    "Domain Group Business Owner",
    "Domain Group Technical Owner",
    "Domain Name",
    "Domain Description",
    "Domain Business Owner",
    "Domain Technical Owner",
    "Subdomain Name",
    "Subdomain Description",
    "Subdomain Business Owner",
    "Subdomain Technical Owner",
    "Data Product Name",
    "Data Product Description",
    "Data Product Business Owner",
    "Data Product Technical Owner",
]

# "Domain Group Technical Owner" is optional (DG tech owner isn't required).
REQUIRED_COLUMNS = [c for c in COLUMNS if c != "Domain Group Technical Owner"]

EXAMPLE_ROW = [
    "Customer", "Customer-facing data assets", "Jane Doe", "",
    "Loyalty", "Loyalty program data", "Jane Doe", "Raj Patel",
    "Membership", "Member profiles and tiers", "Jane Doe", "Raj Patel",
    "Member Profile", "Curated member profile product", "Jane Doe", "Raj Patel",
]


class RowError(Exception):
    pass


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class ExcelService:
    def __init__(self, service: TaxonomyService) -> None:
        self.service = service

    # -- export ----------------------------------------------------------
    def _style(self, ws) -> None:  # noqa: ANN001
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, name in enumerate(COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(18, len(name) + 2)

    def export(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(COLUMNS)

        for g in self.service.get_tree():
            g_cells = [g["name"], g["description"], g["businessOwner"], g["technicalOwner"] or ""]
            if not g["children"]:
                ws.append(g_cells + [""] * 12)
                continue
            for d in g["children"]:
                d_cells = [d["name"], d["description"], d["businessOwner"], d["technicalOwner"] or ""]
                if not d["children"]:
                    ws.append(g_cells + d_cells + [""] * 8)
                    continue
                for s in d["children"]:
                    s_cells = [s["name"], s["description"], s["businessOwner"], s["technicalOwner"] or ""]
                    if not s["children"]:
                        ws.append(g_cells + d_cells + s_cells + [""] * 4)
                        continue
                    for p in s["children"]:
                        ws.append(
                            g_cells + d_cells + s_cells
                            + [p["name"], p["description"], p["businessOwner"], p["technicalOwner"] or ""]
                        )

        self._style(ws)
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def template(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(COLUMNS)
        ws.append(EXAMPLE_ROW)
        self._style(ws)
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # -- import ----------------------------------------------------------
    def import_workbook(self, buffer: bytes, actor: str = "system") -> dict:
        wb = load_workbook(BytesIO(buffer), data_only=True)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        if ws is None:
            raise ValidationError({"file": "Workbook contains no worksheets"})

        header_index: dict[str, int] = {}
        for col, cell in enumerate(ws[1], start=1):
            key = _cell_text(cell.value).lower()
            if key:
                header_index[key] = col

        missing = [c for c in REQUIRED_COLUMNS if c.lower() not in header_index]
        if missing:
            raise ValidationError({"file": f"Missing required column(s): {', '.join(missing)}"})

        def get(row: tuple, name: str) -> str:
            col = header_index.get(name.lower())
            return _cell_text(row[col - 1]) if col else ""

        summary = {"created": 0, "updated": 0, "skipped": 0, "errors": []}

        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            parsed = _parse_row(row, get)
            if not any(parsed.values()):
                summary["skipped"] += 1
                continue
            try:
                _validate_row_shape(parsed)
                for level, payload in (
                    (Level.DOMAIN_GROUP, parsed["dg"]),
                    (Level.DOMAIN, parsed["domain"]),
                    (Level.SUBDOMAIN, parsed["sub"]),
                    (Level.DATA_PRODUCT, parsed["product"]),
                ):
                    if payload:
                        validate_create(level, payload)

                gid = self._upsert(Level.DOMAIN_GROUP, None, parsed["dg"], summary, actor)
                if not parsed["domain"]:
                    continue
                did = self._upsert(Level.DOMAIN, gid, parsed["domain"], summary, actor)
                if not parsed["sub"]:
                    continue
                sid = self._upsert(Level.SUBDOMAIN, did, parsed["sub"], summary, actor)
                if not parsed["product"]:
                    continue
                self._upsert(Level.DATA_PRODUCT, sid, parsed["product"], summary, actor)
            except ValidationError as err:
                summary["errors"].append({"row": row_number, "message": "; ".join(err.fields.values())})
            except RowError as err:
                summary["errors"].append({"row": row_number, "message": str(err)})

        return summary

    def _upsert(self, level: Level, parent_id: Optional[str], payload: dict, summary: dict, actor: str = "system") -> str:
        tech = normalize_technical_owner(payload.get("technicalOwner"))
        name = payload["name"]
        existing = next(
            (n for n in self.service.list(level, parent_id) if same_name(n["name"], name)), None
        )
        if existing:
            changed = (
                existing["description"] != payload["description"].strip()
                or existing["businessOwner"] != payload["businessOwner"].strip()
                or (existing["technicalOwner"] or None) != tech
                or existing["name"] != name.strip()
            )
            if changed:
                self.service.update(
                    level,
                    existing["id"],
                    {
                        "name": name,
                        "description": payload["description"],
                        "businessOwner": payload["businessOwner"],
                        "technicalOwner": tech,
                    },
                    actor=actor,
                )
                summary["updated"] += 1
            return existing["id"]

        created = self.service.create(
            level,
            parent_id,
            {
                "name": name,
                "description": payload["description"],
                "businessOwner": payload["businessOwner"],
                "technicalOwner": tech,
            },
            actor=actor,
        )
        summary["created"] += 1
        return created["id"]


def _parse_row(row: tuple, get) -> dict:  # noqa: ANN001
    def node(n: str, d: str, b: str, t: str) -> Optional[dict]:
        name, description = get(row, n), get(row, d)
        business_owner, technical_owner = get(row, b), get(row, t)
        if not name and not description and not business_owner and not technical_owner:
            return None
        return {
            "name": name,
            "description": description,
            "businessOwner": business_owner,
            "technicalOwner": technical_owner,
        }

    return {
        "dg": node("Domain Group Name", "Domain Group Description", "Domain Group Business Owner", "Domain Group Technical Owner"),
        "domain": node("Domain Name", "Domain Description", "Domain Business Owner", "Domain Technical Owner"),
        "sub": node("Subdomain Name", "Subdomain Description", "Subdomain Business Owner", "Subdomain Technical Owner"),
        "product": node("Data Product Name", "Data Product Description", "Data Product Business Owner", "Data Product Technical Owner"),
    }


def _validate_row_shape(p: dict) -> None:
    if p["product"] and not p["sub"]:
        raise RowError("Data Product present but Subdomain is missing")
    if p["sub"] and not p["domain"]:
        raise RowError("Subdomain present but Domain is missing")
    if p["domain"] and not p["dg"]:
        raise RowError("Domain present but Domain Group is missing")
